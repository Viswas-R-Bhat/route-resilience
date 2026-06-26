# -*- coding: utf-8 -*-
"""
Build the ISRO BAH 2026 Idea Submission deck for Team "Escape the Matrix"
(Problem Statement 4 - Route Resilience), by editing a COPY of the official
template so the branded backgrounds are preserved.
"""
import shutil, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.oxml.ns import qn
from PIL import Image

SRC = "[Pub] ISRO BAH 2026 _ Idea Submission Template.pptx"
OUT = "Escape_the_Matrix_ISRO_BAH_2026_PS4_Idea_Submission.pptx"

# ---------- palette (white-body brand-aligned) ----------
NAVY   = "14233D"   # headings / dark text
INK    = "0F1B2D"
GREY   = "5A6472"   # secondary text
ORANGE = "F26522"   # ISRO accent
ODEEP  = "D44E12"
BLUE   = "2E6FE0"
CYAN   = "1F9BD1"
TEAL   = "12A594"
PURPLE = "7A5C9E"
AMBER  = "E0941A"
RED    = "D63B3B"
LIGHT  = "F3F6FA"   # card fill
CARD2  = "EAF0F7"
BORDER = "D5DCE6"
WHITE  = "FFFFFF"

# dark dashboard mockup palette (the actual product UI)
DBG    = "0B0E14"
DPANEL = "121826"
DCARD  = "1A2233"
DTEAL  = "00E6BD"
DGOLD  = "C9A84C"
DAMBER = "FFB347"
DRED   = "FF4D4D"
DTEXT  = "E8EDF2"
DMUTE  = "8B97A7"
DLINE  = "2A3550"

HFONT = "Segoe UI Semibold"
BFONT = "Segoe UI"
MONO  = "Consolas"

# layout constants (inches)
LEFT, RIGHT = 0.45, 9.55
WIDTH = RIGHT - LEFT
TITLE_TOP = 0.74
CONTENT_TOP = 1.52
BODY_BOT = 5.42

def rgb(h): return RGBColor.from_string(h)

# ---------------- helpers ----------------
def del_nonbg(slide):
    """Remove every shape except the full-bleed background picture."""
    for sh in list(slide.shapes):
        if sh.shape_type == 13:  # PICTURE
            continue
        sh._element.getparent().remove(sh._element)

def _set_runs(tf, lines, default):
    """lines: list of dicts {t, size, color, bold, font, align, space_after, level}"""
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ln.get("align", default.get("align", PP_ALIGN.LEFT))
        if ln.get("space_before") is not None: p.space_before = Pt(ln["space_before"])
        p.space_after = Pt(ln.get("space_after", default.get("space_after", 2)))
        if ln.get("line_spacing"): p.line_spacing = ln["line_spacing"]
        runs = ln["t"] if isinstance(ln["t"], list) else [ln]
        for j, rspec in enumerate(runs):
            r = p.add_run()
            r.text = rspec["t"] if isinstance(rspec, dict) else rspec
            f = r.font
            f.size = Pt(rspec.get("size", ln.get("size", default["size"])) if isinstance(rspec, dict) else ln.get("size", default["size"]))
            f.bold = rspec.get("bold", ln.get("bold", default.get("bold", False))) if isinstance(rspec, dict) else ln.get("bold", default.get("bold", False))
            f.name = rspec.get("font", ln.get("font", default.get("font", BFONT))) if isinstance(rspec, dict) else ln.get("font", default.get("font", BFONT))
            f.color.rgb = rgb(rspec.get("color", ln.get("color", default["color"])) if isinstance(rspec, dict) else ln.get("color", default["color"]))

def add_text(slide, l, t, w, h, lines, size=14, color=NAVY, bold=False, font=BFONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, wrap=True, space_after=2):
    tb = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(tf, m, Emu(0))
    _set_runs(tf, lines, dict(size=size, color=color, bold=bold, font=font, align=align, space_after=space_after))
    return tb

def add_box(slide, l, t, w, h, fill=LIGHT, line=BORDER, line_w=0.75, radius=0.08,
            shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    sp = slide.shapes.add_shape(shape, Inches(l), Inches(t), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = rgb(line); sp.line.width = Pt(line_w)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try: sp.adjustments[0] = radius
        except Exception: pass
    sp.text_frame.word_wrap = True
    for m in ("margin_left", "margin_right", "margin_top", "margin_bottom"):
        setattr(sp.text_frame, m, Emu(0))
    return sp

def box_text(sp, lines, size=12, color=NAVY, bold=False, font=BFONT,
             align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.MIDDLE):
    tf = sp.text_frame
    tf.vertical_anchor = anchor
    _set_runs(tf, lines, dict(size=size, color=color, bold=bold, font=font, align=align, space_after=1))

def add_line(slide, x1, y1, x2, y2, color=GREY, w=1.5, dash=None, arrow=False):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    ln.line.color.rgb = rgb(color); ln.line.width = Pt(w)
    ln.shadow.inherit = False
    lnEl = ln.line._get_or_add_ln()
    if dash:
        d = lnEl.makeelement(qn('a:prstDash'), {'val': dash}); lnEl.append(d)
    if arrow:
        te = lnEl.makeelement(qn('a:tailEnd'), {'type': 'triangle', 'w': 'med', 'len': 'med'})
        lnEl.append(te)
    return ln

def add_oval(slide, cx, cy, r, fill, line=None, line_w=1.0):
    sp = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - r), Inches(cy - r), Inches(2*r), Inches(2*r))
    sp.shadow.inherit = False
    sp.fill.solid(); sp.fill.fore_color.rgb = rgb(fill)
    if line is None: sp.line.fill.background()
    else: sp.line.color.rgb = rgb(line); sp.line.width = Pt(line_w)
    return sp

def content_title(slide, text, sub=None):
    add_box(slide, LEFT, TITLE_TOP + 0.06, 0.16, 0.40, fill=ORANGE, line=None, radius=0.18)
    add_text(slide, LEFT + 0.30, TITLE_TOP, WIDTH - 0.3, 0.55,
             [{"t": text}], size=25, color=NAVY, bold=True, font=HFONT, anchor=MSO_ANCHOR.MIDDLE)
    if sub:
        add_text(slide, LEFT + 0.30, TITLE_TOP + 0.52, WIDTH - 0.3, 0.30,
                 [{"t": sub}], size=11.5, color=GREY, font=BFONT)

# ---------------- build ----------------
shutil.copyfile(SRC, OUT)
prs = Presentation(OUT)
sl = prs.slides

# convert flow webp -> png
Image.open("challenge-4-flow-diagram.webp").convert("RGB").save(".qa/flow.png", quality=95)

# ===== SLIDE 1 : title (fill existing textboxes) =====
s1 = sl[0]
vals = {
    "team name": ("Escape the Matrix", 18),
    "problem statement": ("PS-4  ·  Route Resilience — Occlusion-Robust Road Extraction & Graph-Theoretic Criticality Analysis", 12.5),
    "team leader": ("Viswas R Bhat", 16),
}
for sh in s1.shapes:
    if not sh.has_text_frame: continue
    raw = sh.text_frame.text.strip().lower()
    for key, (val, sz) in vals.items():
        if raw.startswith(key):
            label = sh.text_frame.text.strip()
            if not label.endswith(":"): label = label.rstrip().rstrip(":") + " :"
            tf = sh.text_frame
            tf.clear(); tf.word_wrap = True
            p = tf.paragraphs[0]
            r1 = p.add_run(); r1.text = label + "  "
            r1.font.bold = True; r1.font.size = Pt(sz); r1.font.color.rgb = rgb(NAVY); r1.font.name = HFONT
            r2 = p.add_run(); r2.text = val
            r2.font.bold = False; r2.font.size = Pt(sz); r2.font.color.rgb = rgb(ODEEP if key!="problem statement" else INK); r2.font.name = BFONT
            break

# ===== SLIDE 2 : team table =====
s2 = sl[1]
COLLEGE = "BMS College of Engineering"
members = [
    ("Team Leader", "Viswas R Bhat"),
    ("Team Member-1", "Viswas R Bhat"),
    ("Team Member-2", "Mayank Kumar Singh"),
    ("Team Member-3", "Dibyansh Raj"),
]
tbl = None
for sh in s2.shapes:
    if sh.has_table: tbl = sh.table
cells = [(0,0),(0,1),(1,0),(1,1)]
for (role, name), (r, c) in zip(members, cells):
    cell = tbl.cell(r, c)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf = cell.text_frame; tf.clear(); tf.word_wrap = True
    p0 = tf.paragraphs[0]; run = p0.add_run(); run.text = role
    run.font.bold = True; run.font.size = Pt(13); run.font.color.rgb = rgb(ORANGE); run.font.name = HFONT
    p1 = tf.add_paragraph()
    a = p1.add_run(); a.text = "Name: "; a.font.bold = True; a.font.size = Pt(11.5); a.font.color.rgb = rgb(NAVY); a.font.name = BFONT
    b = p1.add_run(); b.text = name; b.font.size = Pt(11.5); b.font.color.rgb = rgb(INK); b.font.name = BFONT
    p2 = tf.add_paragraph()
    c1 = p2.add_run(); c1.text = "College: "; c1.font.bold = True; c1.font.size = Pt(10.5); c1.font.color.rgb = rgb(NAVY); c1.font.name = BFONT
    c2 = p2.add_run(); c2.text = COLLEGE; c2.font.size = Pt(10.5); c2.font.color.rgb = rgb(GREY); c2.font.name = BFONT

# ===== SLIDE 3 : Opportunity & USP =====
s3 = sl[2]; del_nonbg(s3)
content_title(s3, "Opportunity & Unique Selling Proposition")
# problem line
add_text(s3, LEFT, CONTENT_TOP - 0.06, WIDTH, 0.62,
    [{"t":[{"t":"The gap:  ","bold":True,"color":NAVY},
           {"t":"Satellite road extraction suffers “spectral blindness” — tree canopy, building shadows and clouds break the mask, leaving it topologically disconnected and useless for routing, traffic or disaster response.","color":GREY}],
      "size":12.5,"line_spacing":1.05}], size=12.5)
# left column: how different + USP
lx, lw = LEFT, 4.55
add_text(s3, lx, CONTENT_TOP+0.60, lw, 0.3, [{"t":"How we are different"}], size=13.5, color=BLUE, bold=True, font=HFONT)
usps = [
    ("Occlusion-robust extraction", "Canopy-aware hysteresis grows faint roads from confident seeds, reaching deeper under tree cover to recover occluded segments.", TEAL),
    ("Evidence-based topological healing", "Union-Find bridges gaps — endpoint-pairs AND mid-edge T-junctions — only where evidence agrees the road continues: the model's sub-threshold confidence, an Excess-Green canopy mask and trajectory alignment, each tagged high/med/low.", BLUE),
    ("Predictive resilience: failure & flood", "Weighted-betweenness ablation, DEM flood-submersion and demand-weighted criticality quantify systemic collapse.", ORANGE),
    ("Validated against OSM ground truth", "A topological-accuracy benchmark scores the graph vs OpenStreetMap: recall, path-length error and routing match.", PURPLE),
]
cy = CONTENT_TOP + 0.98
ch = 0.66
for title, desc, col in usps:
    add_box(s3, lx, cy, 0.10, ch-0.12, fill=col, line=None, radius=0.3)
    add_text(s3, lx+0.22, cy-0.02, lw-0.22, ch,
        [{"t":title,"bold":True,"size":12,"color":NAVY,"space_after":1},
         {"t":desc,"size":10.0,"color":GREY,"line_spacing":1.0}], size=11)
    cy += ch + 0.05
# right column: before/after graphic
rx = 5.35; rw = RIGHT - rx
add_text(s3, rx, CONTENT_TOP+0.60, rw, 0.3, [{"t":"Broken mask  →  Healed routable network"}], size=13.5, color=BLUE, bold=True, font=HFONT)
panelY = CONTENT_TOP + 0.98; panelH = 2.28; panelW = (rw-0.5)/2
# panel A - broken
pa = add_box(s3, rx, panelY, panelW, panelH, fill=LIGHT, line=BORDER, radius=0.06)
add_text(s3, rx, panelY+panelH+0.03, panelW, 0.25, [{"t":"Standard extraction","align":PP_ALIGN.CENTER}], size=9.5, color=RED, bold=True)
ax, ay, aw, ah = rx+0.12, panelY+0.12, panelW-0.24, panelH-0.24
seg = [((0.05,0.2),(0.4,0.25)),((0.55,0.3),(0.95,0.35)),((0.2,0.55),(0.45,0.85)),((0.6,0.6),(0.6,0.95)),((0.05,0.8),(0.3,0.82)),((0.7,0.15),(0.72,0.5))]
for (x1f,y1f),(x2f,y2f) in seg:
    add_line(s3, ax+aw*x1f, ay+ah*y1f, ax+aw*x2f, ay+ah*y2f, color=RED, w=2.0, dash="dash")
# panel B - healed
pbx = rx+panelW+0.5
pb = add_box(s3, pbx, panelY, panelW, panelH, fill=LIGHT, line=BORDER, radius=0.06)
add_text(s3, pbx, panelY+panelH+0.03, panelW, 0.25, [{"t":"Route Resilience","align":PP_ALIGN.CENTER}], size=9.5, color=TEAL, bold=True)
bx, by, bw, bh = pbx+0.12, panelY+0.12, panelW-0.24, panelH-0.24
nodes = {"a":(0.15,0.25),"b":(0.5,0.2),"c":(0.85,0.3),"d":(0.25,0.6),"e":(0.55,0.62),"f":(0.8,0.8),"g":(0.2,0.9)}
edges = [("a","b"),("b","c"),("a","d"),("b","e"),("c","f"),("d","e"),("e","f"),("d","g"),("e","g")]
for u,v in edges:
    x1f,y1f = nodes[u]; x2f,y2f = nodes[v]
    add_line(s3, bx+bw*x1f, by+bh*y1f, bx+bw*x2f, by+bh*y2f, color=TEAL, w=2.0)
for k,(xf,yf) in nodes.items():
    col = RED if k=="e" else BLUE
    rr = 0.058 if k=="e" else 0.040
    add_oval(s3, bx+bw*xf, by+bh*yf, rr, fill=col, line=WHITE, line_w=1.0)
# arrow between panels
add_line(s3, rx+panelW+0.06, panelY+panelH/2, pbx-0.06, panelY+panelH/2, color=ORANGE, w=2.5, arrow=True)

# ===== SLIDE 4 : Features =====
s4 = sl[3]; del_nonbg(s4)
content_title(s4, "Key Features")
feats = [
    ("Occlusion-robust segmentation", "U-Net/ResNet34 with canopy-aware hysteresis thresholding recovers roads under canopy & shadow.", TEAL),
    ("Topological reconstruction", "Skeletonize → graph → Union-Find healing yields a routable weighted vector network.", BLUE),
    ("Evidence-based gap healing", "Endpoint + mid-edge T-junction bridges gated by model soft-confidence, Excess-Green canopy & trajectory — each tagged high/med/low.", TEAL),
    ("Gatekeeper-node detection", "Weighted betweenness (k-sampled at scale) maps critical intersections & weakest links.", ORANGE),
    ("Flood & failure stress-test", "DEM-driven flood submersion + node ablation drive a live Resilience Index.", RED),
    ("Demand-weighted criticality", "Betweenness × local road-density surfaces nodes whose failure strands real travel.", AMBER),
    ("OSM validation benchmark", "Graph scored vs OpenStreetMap: recall, path-length error & routing match.", PURPLE),
    ("Live interactive dashboard", "Custom Leaflet + Three.js (WebGL) UI — search any place for live GPU analysis.", CYAN),
]
cols, rows = 2, 4
gx, gy = 0.30, 0.18
cw = (WIDTH - gx) / cols
chh = (BODY_BOT - CONTENT_TOP - gy*(rows-1)) / rows
for idx, (title, desc, col) in enumerate(feats):
    r, c = divmod(idx, cols)
    x = LEFT + c*(cw+gx); y = CONTENT_TOP + r*(chh+gy)
    add_box(s4, x, y, cw, chh, fill=LIGHT, line=BORDER, radius=0.07)
    add_oval(s4, x+0.34, y+chh/2, 0.18, fill=col, line=None)
    add_text(s4, x+0.30, y+chh/2-0.16, 0.36, 0.34, [{"t":str(idx+1),"align":PP_ALIGN.CENTER}], size=12, color=WHITE, bold=True, font=HFONT, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s4, x+0.64, y+0.10, cw-0.78, chh-0.16,
        [{"t":title,"bold":True,"size":12,"color":NAVY,"space_after":1},
         {"t":desc,"size":9.8,"color":GREY,"line_spacing":1.0}], anchor=MSO_ANCHOR.MIDDLE)

# ===== SLIDE 5 : Process flow (embed image, framed) =====
s5 = sl[4]; del_nonbg(s5)
content_title(s5, "Process Flow",
    sub="Imagery → extraction → skeleton/graph → canopy-aware healing → resilience → OSM validation → dashboard")
iw, ih = Image.open(".qa/flow.png").size
aspect = ih/iw
region_top = CONTENT_TOP + 0.36
region_bot = BODY_BOT
max_h = (region_bot - region_top) - 0.10
disp_h = min(max_h, WIDTH * aspect)
disp_w = disp_h / aspect
pad = 0.12
fw, fh = disp_w + 2*pad, disp_h + 2*pad
fx = LEFT + (WIDTH - fw)/2
fy = region_top + ((region_bot - region_top) - fh)/2
add_box(s5, fx, fy, fw, fh, fill=WHITE, line=BORDER, line_w=1.0, radius=0.03)
s5.shapes.add_picture(".qa/flow.png", Inches(fx+pad), Inches(fy+pad), width=Inches(disp_w), height=Inches(disp_h))

# ===== SLIDE 6 : dashboard mockup (dark) =====
s6 = sl[5]; del_nonbg(s6)
content_title(s6, "Live Dashboard", sub="Premium dark UI · Leaflet + Three.js (WebGL) · live GPU backend · search any place")
mY = CONTENT_TOP + 0.30
app = add_box(s6, LEFT, mY, WIDTH, BODY_BOT-mY, fill=DBG, line=DLINE, radius=0.03)
# app header
add_text(s6, LEFT+0.25, mY+0.12, 4.5, 0.3, [{"t":"◉  ROUTE RESILIENCE","bold":True}], size=14, color=DTEAL, font=HFONT)
add_text(s6, RIGHT-3.0, mY+0.16, 2.8, 0.3, [{"t":"ISRO BAH 2026 · PS-4","align":PP_ALIGN.RIGHT}], size=9, color=DMUTE)
# metric cards row
metrics = [("RESILIENCE INDEX","0.78",DTEAL,"HSR Layout · stable"),("ROAD NODES","187",DTEAL,"healed graph"),
           ("CONNECTIVITY","90%",DTEAL,"LCC post-healing"),("OSM RECALL","65%",DGOLD,"vs ground truth")]
mcw = (WIDTH-0.5-0.30*3)/4; mcx = LEFT+0.25; mcy = mY+0.52; mch=0.80
for label, val, col, sub in metrics:
    add_box(s6, mcx, mcy, mcw, mch, fill=DCARD, line=DLINE, radius=0.10)
    add_text(s6, mcx+0.12, mcy+0.08, mcw-0.2, 0.2, [{"t":label}], size=7.5, color=DMUTE, bold=True)
    add_text(s6, mcx+0.12, mcy+0.26, mcw-0.2, 0.34, [{"t":val}], size=19, color=col, bold=True, font=MONO)
    add_text(s6, mcx+0.12, mcy+0.60, mcw-0.2, 0.18, [{"t":sub}], size=7.5, color=col)
    mcx += mcw+0.30
# map panel
mapY = mcy+mch+0.16; mapH = BODY_BOT-0.18-mapY; mapW = WIDTH-0.5-3.0-0.25
mapX = LEFT+0.25
mp = add_box(s6, mapX, mapY, mapW, mapH, fill="0E1320", line=DLINE, radius=0.04)
add_text(s6, mapX+0.12, mapY+0.06, 4, 0.2, [{"t":"CRITICALITY MAP · HSR Layout, Bengaluru"}], size=8, color=DMUTE, bold=True)
gx0, gy0, gw0, gh0 = mapX+0.2, mapY+0.34, mapW-0.4, mapH-0.5
mnodes = {"a":(0.1,0.2),"b":(0.4,0.12),"c":(0.72,0.22),"d":(0.92,0.4),"e":(0.2,0.5),
          "f":(0.5,0.45),"g":(0.78,0.55),"h":(0.32,0.82),"i":(0.62,0.8),"j":(0.88,0.85)}
medges=[("a","b"),("b","c"),("c","d"),("a","e"),("b","f"),("c","g"),("e","f"),("f","g"),
        ("g","d"),("e","h"),("f","i"),("h","i"),("i","j"),("g","j")]
emph=[("c","d"),("f","g"),("c","g")]  # critical edges in amber/red
for u,v in medges:
    x1f,y1f=mnodes[u]; x2f,y2f=mnodes[v]
    if (u,v) in emph: col,w = DRED,3.0
    elif u in ("c","f","g") or v in ("c","f","g"): col,w = DAMBER,2.4
    else: col,w = DTEAL,1.6
    add_line(s6, gx0+gw0*x1f, gy0+gh0*y1f, gx0+gw0*x2f, gy0+gh0*y2f, color=col, w=w)
for k,(xf,yf) in mnodes.items():
    if k=="f": col,rr=DRED,0.060
    elif k in ("c","g"): col,rr=DAMBER,0.046
    else: col,rr=DTEAL,0.034
    add_oval(s6, gx0+gw0*xf, gy0+gh0*yf, rr, fill=col, line=DBG, line_w=0.75)
# ablated node "x" mark on node f
fx,fy = gx0+gw0*mnodes["f"][0], gy0+gh0*mnodes["f"][1]
add_text(s6, fx-0.12, fy-0.14, 0.24, 0.26, [{"t":"✕","align":PP_ALIGN.CENTER}], size=12, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
# right side: gatekeeper list + legend
gkx = mapX+mapW+0.25; gkw = RIGHT-0.25-gkx
add_box(s6, gkx, mapY, gkw, mapH, fill=DPANEL, line=DLINE, radius=0.06)
add_text(s6, gkx+0.14, mapY+0.08, gkw-0.28, 0.2, [{"t":"GATEKEEPER NODES"}], size=8, color=DGOLD, bold=True)
ranks=[("#1  N-78","0.390",DRED),("#2  N-73","0.312",DRED),("#3  N-74","0.284",DAMBER),
       ("#4  N-69","0.267",DAMBER),("#5  N-76","0.265",DTEAL)]
ry=mapY+0.32; row_h=0.26; step=0.295
for name,sc,col in ranks:
    add_box(s6, gkx+0.12, ry, gkw-0.24, row_h, fill=DCARD, line=None, radius=0.20)
    add_text(s6, gkx+0.24, ry, gkw-0.9, row_h, [{"t":name}], size=8.5, color=DTEXT, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    add_text(s6, gkx+gkw-0.72, ry, 0.6, row_h, [{"t":sc,"align":PP_ALIGN.RIGHT}], size=8.5, color=col, font=MONO, anchor=MSO_ANCHOR.MIDDLE)
    ry+=step

# ===== SLIDE 7 : Architecture =====
s7 = sl[6]; del_nonbg(s7)
content_title(s7, "Solution Architecture")
# input strip
inY = CONTENT_TOP + 0.02
ib = add_box(s7, LEFT, inY, WIDTH, 0.46, fill=NAVY, line=None, radius=0.10)
box_text(ib, [{"t":[{"t":"INPUTS   ","bold":True,"color":ORANGE,"size":11},
                    {"t":"DeepGlobe (training)  ·  Esri World Imagery ~0.6 m (live)  ·  AWS DEM  ·  OpenStreetMap (validation)  ·  NIR-ready for LISS-IV / Cartosat-3","color":WHITE,"size":9.8}]}],
         align=PP_ALIGN.CENTER)
# phase boxes
phases = [
    ("P1","Segmentation","U-Net + ResNet34 (SMP)\nCanopy-aware hysteresis\nD4 TTA · occlusion aug", TEAL),
    ("P2","Skeleton → Graph","Zhang-Suen thinning\nsknw graph + spur prune\nlength-weighted edges", BLUE),
    ("P3","Canopy Healing","Union-Find + Excess-Green\nTrajectory alignment\nConfidence-tiered bridges", PURPLE),
    ("P4","Resilience","Betweenness (k-sampled)\nFlood-DEM + ablation\nDemand-weighted crit.", ORANGE),
    ("P5","Live Dashboard","Leaflet + Three.js (WebGL)\nECharts · /api/analyze\nstdlib Python server", CYAN),
]
art_out = ["road mask","raw graph","routable graph","criticality + RI","live planner UI"]
n=5; gap=0.34
pbw=(WIDTH-gap*(n-1))/n
pbY=inY+0.78; pbH=2.05
xs=[]
for i,(num,name,body,col) in enumerate(phases):
    x=LEFT+i*(pbw+gap); xs.append(x)
    add_box(s7, x, pbY, pbw, pbH, fill=WHITE, line=col, line_w=1.75, radius=0.06)
    hdr=add_box(s7, x, pbY, pbw, 0.50, fill=col, line=None, radius=0.06)
    box_text(hdr, [{"t":[{"t":num+"  ","bold":True,"size":12,"color":WHITE},{"t":name,"bold":True,"size":10.5,"color":WHITE}]}], align=PP_ALIGN.CENTER)
    add_text(s7, x+0.10, pbY+0.60, pbw-0.20, 0.95,
             [{"t":ln,"size":8.6,"color":NAVY,"align":PP_ALIGN.CENTER,"space_after":2,"line_spacing":1.02} for ln in body.split("\n")],
             anchor=MSO_ANCHOR.TOP)
    # divider + output-artifact footer inside the box (avoids cramped between-box labels)
    add_line(s7, x+0.18, pbY+pbH-0.46, x+pbw-0.18, pbY+pbH-0.46, color=BORDER, w=1.0)
    add_text(s7, x+0.08, pbY+pbH-0.42, pbw-0.16, 0.34,
             [{"t":[{"t":"▸ ","bold":True,"color":col},{"t":art_out[i],"color":col,"bold":True}],"align":PP_ALIGN.CENTER}],
             size=8.4, anchor=MSO_ANCHOR.MIDDLE)
# arrows only, mid-height (data flows left → right)
for i in range(n-1):
    x1=xs[i]+pbw; x2=xs[i+1]; midy=pbY+pbH/2
    add_line(s7, x1+0.02, midy, x2-0.02, midy, color=GREY, w=2.25, arrow=True)
# outcome strip
outY=pbY+pbH+0.16
ob=add_box(s7, LEFT, outY, WIDTH, 0.40, fill=LIGHT, line=BORDER, radius=0.10)
box_text(ob, [{"t":[{"t":"VALIDATED   ","bold":True,"color":TEAL,"size":10.5},
                    {"t":"Held-out DeepGlobe (934 tiles): IoU 0.605 · Relaxed-IoU 0.770   ·   vs OpenStreetMap: recall 65%, median path-error 7.5%   ·   healing LCC 77%→90%","color":NAVY,"size":9.0}]}],
         align=PP_ALIGN.CENTER)

# ===== SLIDE 8 : Technologies =====
s8 = sl[7]; del_nonbg(s8)
content_title(s8, "Technology Stack")

def chip_row(slide, chips, x0, x_right, y_center, base=9.5, h=0.32, gap=0.11):
    cwf = lambda t, fs: 0.30 + len(t)*0.0082*fs
    widths = [cwf(c, base) for c in chips]
    total = sum(widths) + gap*(len(chips)-1)
    avail = x_right - x0
    fs = base
    if total > avail:  # scale font + widths so the row fits on ONE line (no overlap)
        fs = max(7.0, base * (avail/total))
        widths = [cwf(c, fs) for c in chips]
        if sum(widths) + gap*(len(chips)-1) > avail and len(chips) > 1:
            gap = max(0.05, (avail - sum(widths))/(len(chips)-1))
    x = x0
    for c, w in zip(chips, widths):
        add_box(slide, x, y_center-h/2, w, h, fill=CARD2, line=BORDER, radius=0.5)
        add_text(slide, x+0.05, y_center-h/2, w-0.10, h, [{"t":c,"align":PP_ALIGN.CENTER}],
                 size=fs, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
        x += w + gap

groups = [
    ("Data & Imagery", ORANGE, ["NumPy","OpenCV","Pillow","Albumentations","Web-Mercator tiling"]),
    ("Segmentation (DL)", TEAL, ["PyTorch","SMP · U-Net/ResNet34","D4 TTA","Mixed-Precision (AMP)","clDice loss"]),
    ("Skeleton & Graph", BLUE, ["scikit-image · Zhang-Suen","sknw","NetworkX","SciPy cKDTree"]),
    ("Resilience & Analysis", PURPLE, ["Betweenness (k-sampled)","Union-Find healing","Flood-DEM ablation","Demand weighting"]),
    ("Dashboard & Backend", CYAN, ["Vanilla JS","Leaflet.js","Three.js (WebGL)","Apache ECharts","Python http.server"]),
    ("Data Sources", AMBER, ["Esri World Imagery","OpenStreetMap / Overpass","AWS DEM","DeepGlobe","LISS-IV / Cartosat (NIR-ready)"]),
]
rows = len(groups); rgap = 0.14
rh = (BODY_BOT - CONTENT_TOP - rgap*(rows-1)) / rows
ry = CONTENT_TOP; tagw = 1.95
for name, col, chips in groups:
    add_box(s8, LEFT, ry, tagw, rh, fill=col, line=None, radius=0.10)
    add_text(s8, LEFT+0.14, ry, tagw-0.22, rh, [{"t":name}], size=10.5, color=WHITE, bold=True, font=HFONT, anchor=MSO_ANCHOR.MIDDLE)
    chip_row(s8, chips, LEFT+tagw+0.22, RIGHT, ry + rh/2)
    ry += rh + rgap

# ===== SLIDE 9 : Estimated cost =====
s9 = sl[8]; del_nonbg(s9)
content_title(s9, "Honest Cost Breakdown", sub="Open data & open source — the only real recurring cost is the GPU that powers live inference")
rows = [
    ("Satellite & map imagery", "Esri tiles · OpenStreetMap / Overpass · AWS DEM — open, no API key", "₹0", TEAL),
    ("Training & validation data", "DeepGlobe (training) · OpenStreetMap (validation) — open", "₹0", TEAL),
    ("Model training", "One-time, a few GPU-hours on the team's own laptop GPU", "₹0", TEAL),
    ("Live inference — GPU backend", "≈ 50 s / scene; cloud GPU ≈ ₹1–2 per scene, or ₹0 on an owned GPU", "≈₹1–2", AMBER),
    ("Static dashboard hosting", "Vanilla-JS bundle on Vercel / Netlify free tier — no server", "₹0", TEAL),
]
ry = CONTENT_TOP + 0.05
rh = 0.52; rgap = 0.115
for label, detail, cost, col in rows:
    add_box(s9, LEFT, ry, WIDTH-1.7, rh, fill=LIGHT, line=BORDER, radius=0.08)
    add_text(s9, LEFT+0.22, ry+0.07, WIDTH-1.7-2.6, rh-0.12,
             [{"t":label,"bold":True,"size":12,"color":NAVY,"space_after":1},
              {"t":detail,"size":9.3,"color":GREY}], anchor=MSO_ANCHOR.MIDDLE)
    cb = add_box(s9, RIGHT-1.55, ry, 1.55, rh, fill=col, line=None, radius=0.10)
    box_text(cb, [{"t":cost}], size=(18 if len(cost) <= 3 else 15), color=WHITE, bold=True, font=HFONT, align=PP_ALIGN.CENTER)
    ry += rh + rgap
# honest bottom line
tb = add_box(s9, LEFT, ry+0.02, WIDTH, BODY_BOT-(ry+0.02), fill=NAVY, line=None, radius=0.06)
box_text(tb, [
    {"t":[{"t":"DEV & DEMO   ","bold":True,"size":15,"color":WHITE},
          {"t":"≈  ₹0","bold":True,"size":18,"color":ORANGE},
          {"t":"   ·   open data · owned GPU · free static hosting","size":11,"color":CARD2}],
     "align":PP_ALIGN.CENTER,"space_after":3},
    {"t":"At scale the only recurring cost is the GPU for live inference — ≈ ₹1–2 per analysed scene on a cloud GPU, ₹0 on owned hardware.",
     "size":10.5,"color":CARD2,"align":PP_ALIGN.CENTER}], anchor=MSO_ANCHOR.MIDDLE)

prs.save(OUT)
print("Saved", OUT, "with", len(prs.slides.__iter__.__self__._sldIdLst), "slides")
